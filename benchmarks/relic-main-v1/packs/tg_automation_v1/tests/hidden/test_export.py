"""Tests for tg_automation.export.export_report."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from tg_automation.export import export_report, _normalize_field


def test_csv_export_headers_and_rows(tmp_path: Path):
    rows = [
        {"user_id": 1, "messages_sent": 3, "last_active_at": "2026-05-01"},
        {"user_id": 2, "messages_sent": 1, "last_active_at": "2026-05-02"},
    ]
    out = tmp_path / "report.csv"
    export_report(rows, out)
    with open(out, encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        assert reader.fieldnames == ["user_id", "messages_sent", "last_active_at"]
        data = list(reader)
    assert len(data) == 2
    assert data[0]["user_id"] == "1"


def test_xlsx_export_writes_workbook(tmp_path: Path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    rows = [{"a_field": 1, "b_field": "x"}, {"a_field": 2, "b_field": "y"}]
    out = tmp_path / "report.xlsx"
    export_report(rows, out)
    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == ["a_field", "b_field"]
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert body == [[1, "x"], [2, "y"]]


def test_unsupported_extension_raises(tmp_path: Path):
    with pytest.raises(ValueError):
        export_report([{"a": 1}], tmp_path / "x.txt")


def test_headers_are_snake_case():
    assert _normalize_field("Messages Sent") == "messages_sent"
    assert _normalize_field("Reactions-Given") == "reactions_given"
    assert _normalize_field("  active_days  ") == "active_days"


def test_headers_normalized_in_export(tmp_path: Path):
    rows = [{"User ID": 1, "Messages Sent": 3}]
    out = tmp_path / "norm.csv"
    export_report(rows, out)
    with open(out, encoding="utf-8") as fh:
        header = fh.readline().strip()
    assert header == "user_id,messages_sent"


def test_accepts_dataclass_with_as_dict(tmp_path: Path):
    from tg_automation.analysis import UserInteraction
    rec = UserInteraction(user_id=1, username="a", messages_sent=2,
                          reactions_given=1, active_days=1)
    out = tmp_path / "rec.csv"
    export_report([rec], out)
    text = out.read_text(encoding="utf-8")
    assert "user_id" in text
    assert ",1" in text or "1," in text
