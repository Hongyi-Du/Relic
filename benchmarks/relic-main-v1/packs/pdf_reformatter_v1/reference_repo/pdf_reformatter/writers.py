"""Emit reformatted output as PDF, DOCX, and CSV.

Each writer takes the already-paginated rows (or raw rows for the
formats that do not paginate) and produces a file at the requested
path. Writers never re-measure or re-paginate — that responsibility
lives in ``pdf_reformatter.layout``.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import List, Sequence

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch  # noqa: F401  (kept for caller convenience)
from reportlab.pdfgen import canvas

from pdf_reformatter.extract import Row
from pdf_reformatter.layout import (
    DEFAULT_BOTTOM_MARGIN,
    DEFAULT_FONT_SIZE,
    DEFAULT_LEADING,
    DEFAULT_PAGE_HEIGHT,
    DEFAULT_PAGE_WIDTH,
    DEFAULT_TOP_MARGIN,
    Page,
    _wrap_cell,
)


def write_pdf(
    pages: Sequence[Page],
    output_path: str | Path,
    col_widths: Sequence[float],
    page_width: float = DEFAULT_PAGE_WIDTH,
    page_height: float = DEFAULT_PAGE_HEIGHT,
    top_margin: float = DEFAULT_TOP_MARGIN,
    bottom_margin: float = DEFAULT_BOTTOM_MARGIN,
    font_size: float = DEFAULT_FONT_SIZE,
    leading: float = DEFAULT_LEADING,
    avg_char_width: float = 5.0,
) -> Path:
    """Write paginated rows to a PDF file.

    Each ``Page`` becomes one PDF page. Rows are drawn with the same
    column widths used to paginate them so the rendered geometry
    matches what ``paginate`` measured. A box is drawn around every
    cell to keep the table appearance of the source PDF.
    """

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(output), pagesize=(page_width, page_height))
    left = (page_width - sum(col_widths)) / 2.0
    if left < 0:
        left = 0.0

    for page in pages:
        y = page_height - top_margin
        c.setFont("Helvetica", font_size)
        for row, row_height in zip(page.rows, page.heights):
            x = left
            for cell, width in zip(row.cells, col_widths):
                # Cell box.
                c.rect(x, y - row_height, width, row_height, stroke=1, fill=0)
                # Wrapped cell text, top-aligned.
                lines = _wrap_cell(cell, width, avg_char_width)
                ty = y - leading + (leading - font_size) / 2.0
                for line in lines:
                    c.drawString(x + 2.0, ty, line)
                    ty -= leading
                x += width
            y -= row_height
        c.showPage()

    c.save()
    return output


def write_docx(
    rows: Sequence[Row],
    output_path: str | Path,
) -> Path:
    """Write rows to a DOCX file as a single table.

    Column order matches the source; empty cells are preserved as empty
    string cells (not skipped) so the DOCX schema is row-cell-uniform.
    """

    if not rows:
        # Still produce a valid DOCX file so downstream tooling has a
        # consistent contract.
        doc = Document()
        doc.add_paragraph("")
        doc.save(str(output_path))
        return Path(output_path)

    width = max(len(r.cells) for r in rows)
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        cells = row.cells + [""] * (width - len(row.cells))
        for c_idx, value in enumerate(cells):
            table.rows[r_idx].cells[c_idx].text = value
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def write_csv(
    rows: Sequence[Row],
    output_path: str | Path,
) -> Path:
    """Write rows to a CSV file using the stdlib ``csv`` module.

    Empty cells are written as empty fields. The CSV is round-trippable
    by ``csv.reader`` without errors.
    """

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        output.write_text("")
        return output
    width = max(len(r.cells) for r in rows)
    with output.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        for row in rows:
            cells: List[str] = list(row.cells)
            if len(cells) < width:
                cells = cells + [""] * (width - len(cells))
            writer.writerow(cells)
    return output
