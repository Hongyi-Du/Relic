"""End-to-end CLI: reformat sample PDF → PDF + DOCX + CSV."""

from __future__ import annotations

import csv as csv_mod
import subprocess
import sys
from pathlib import Path

import pypdf
import pytest
from docx import Document

from pdf_reformatter.cli import main, reformat


def test_reformat_produces_all_three_outputs(sample_pdf, out_dir):
    result = reformat(
        sample_pdf,
        out_dir / "out.pdf",
        out_dir / "out.docx",
        out_dir / "out.csv",
    )
    assert result["pdf"].exists()
    assert result["docx"].exists()
    assert result["csv"].exists()
    assert result["rows"] > 0
    assert result["pages"] >= 1


def test_reformatted_pdf_no_overflow(sample_pdf, out_dir):
    # Force smaller page so we observe multi-page output.
    result = reformat(
        sample_pdf,
        out_dir / "out.pdf",
        out_dir / "out.docx",
        out_dir / "out.csv",
        page_height=300.0,
    )
    assert result["pages"] >= 2
    reader = pypdf.PdfReader(str(result["pdf"]))
    assert len(reader.pages) == result["pages"]


def test_reformat_csv_matches_extraction(sample_pdf, out_dir):
    result = reformat(
        sample_pdf,
        out_dir / "out.pdf",
        out_dir / "out.docx",
        out_dir / "out.csv",
    )
    with result["csv"].open(newline="") as fh:
        loaded = list(csv_mod.reader(fh))
    # Every row has the same column count and at least one long-note row
    # was preserved end-to-end.
    widths = {len(r) for r in loaded}
    assert len(widths) == 1
    assert any("Very long note" in cell for row in loaded for cell in row)


def test_reformat_docx_matches_csv(sample_pdf, out_dir):
    result = reformat(
        sample_pdf,
        out_dir / "out.pdf",
        out_dir / "out.docx",
        out_dir / "out.csv",
    )
    with result["csv"].open(newline="") as fh:
        csv_rows = list(csv_mod.reader(fh))
    doc = Document(str(result["docx"]))
    table = doc.tables[0]
    assert len(table.rows) == len(csv_rows)
    for r_idx, csv_row in enumerate(csv_rows):
        for c_idx, value in enumerate(csv_row):
            assert table.rows[r_idx].cells[c_idx].text == value


def test_cli_main_runs(sample_pdf, out_dir, capsys):
    rc = main(
        [
            str(sample_pdf),
            "--out-pdf",
            str(out_dir / "cli.pdf"),
            "--out-docx",
            str(out_dir / "cli.docx"),
            "--out-csv",
            str(out_dir / "cli.csv"),
        ]
    )
    assert rc == 0
    captured = capsys.readouterr()
    assert "Wrote" in captured.out
    assert (out_dir / "cli.pdf").exists()
    assert (out_dir / "cli.docx").exists()
    assert (out_dir / "cli.csv").exists()


def test_cli_module_invocation(sample_pdf, out_dir):
    """python -m pdf_reformatter must work as a script entry point."""

    env = {**__import__("os").environ}
    solution_dir = Path(__file__).resolve().parent.parent / "solution"
    if solution_dir.exists():
        env["PYTHONPATH"] = str(solution_dir) + ":" + env.get("PYTHONPATH", "")
    else:
        env["PYTHONPATH"] = "/solution:" + env.get("PYTHONPATH", "")
    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "pdf_reformatter",
            str(sample_pdf),
            "--out-pdf",
            str(out_dir / "m.pdf"),
            "--out-docx",
            str(out_dir / "m.docx"),
            "--out-csv",
            str(out_dir / "m.csv"),
        ],
        env=env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr



def test_reformat_repeat_header_pdf_only_not_docx_csv(sample_pdf, out_dir):
    result = reformat(
        sample_pdf,
        out_dir / "repeat.pdf",
        out_dir / "repeat.docx",
        out_dir / "repeat.csv",
        page_height=300.0,
        repeat_header=True,
    )
    assert result["pages"] >= 2

    with result["csv"].open(newline="") as fh:
        rows = list(csv_mod.reader(fh))
    header = rows[0]
    assert rows.count(header) == 1

    doc = Document(str(result["docx"]))
    table_rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert table_rows.count(header) == 1
