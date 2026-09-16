"""PDF/DOCX/CSV writers."""

from __future__ import annotations

import csv as csv_mod
from pathlib import Path

import pypdf
from docx import Document

from pdf_reformatter.extract import Row, rows_from_records
from pdf_reformatter.layout import paginate
from pdf_reformatter.writers import write_csv, write_docx, write_pdf


COL_WIDTHS = [60, 120, 70, 250]


def _sample_rows() -> list[Row]:
    records = [["ID", "Name", "Qty", "Notes"]]
    for i in range(1, 11):
        notes = "long " * 20 if i == 4 else ""
        records.append([f"{i}", f"Item {i}", f"{i*2}", notes])
    return rows_from_records(records)


def test_write_csv_matches_row_grid(tmp_path):
    rows = _sample_rows()
    path = write_csv(rows, tmp_path / "out.csv")
    with path.open(newline="") as fh:
        loaded = list(csv_mod.reader(fh))
    assert loaded[0] == rows[0].cells
    assert len(loaded) == len(rows)
    assert all(len(r) == len(rows[0].cells) for r in loaded)


def test_write_csv_empty_cells_kept_as_fields(tmp_path):
    rows = rows_from_records([["a", "", "c"], ["", "", ""]])
    path = write_csv(rows, tmp_path / "e.csv")
    with path.open(newline="") as fh:
        loaded = list(csv_mod.reader(fh))
    assert loaded == [["a", "", "c"], ["", "", ""]]


def test_write_docx_round_trips_rows(tmp_path):
    rows = _sample_rows()
    path = write_docx(rows, tmp_path / "out.docx")
    doc = Document(str(path))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == len(rows)
    assert len(table.columns) == len(rows[0].cells)
    # Header cell-for-cell match.
    for c_idx, header in enumerate(rows[0].cells):
        assert table.rows[0].cells[c_idx].text == header


def test_write_pdf_pages_do_not_overflow(tmp_path):
    rows = _sample_rows()
    # Force multi-page output.
    page_height = 250.0
    pages = paginate(
        rows,
        col_widths=COL_WIDTHS,
        page_height=page_height,
    )
    path = write_pdf(pages, tmp_path / "out.pdf", COL_WIDTHS, page_height=page_height)
    reader = pypdf.PdfReader(str(path))
    assert len(reader.pages) == len(pages)
    # Every PDF page exists and has a media box matching our page size.
    for pdf_page in reader.pages:
        height = float(pdf_page.mediabox.height)
        assert abs(height - page_height) < 1.0


def test_write_pdf_emits_at_least_one_page_for_single_row(tmp_path):
    rows = rows_from_records([["solo", "row", "test", ""]])
    pages = paginate(rows, col_widths=COL_WIDTHS)
    path = write_pdf(pages, tmp_path / "solo.pdf", COL_WIDTHS)
    reader = pypdf.PdfReader(str(path))
    assert len(reader.pages) == 1


def test_write_csv_parseable_by_stdlib(tmp_path):
    rows = _sample_rows()
    path = write_csv(rows, tmp_path / "stdlib.csv")
    # Must not raise.
    with path.open(newline="") as fh:
        list(csv_mod.reader(fh))



def test_docx_and_csv_do_not_duplicate_repeated_pdf_headers(tmp_path):
    rows = rows_from_records(
        [["ID", "Name", "Qty", "Notes"]]
        + [[f"{i}", "Item", "1", "note " * 12] for i in range(1, 8)]
    )
    pages = paginate(
        rows,
        col_widths=COL_WIDTHS,
        page_height=170.0,
        repeat_header=True,
    )
    assert any(p.repeated_header for p in pages[1:])

    pdf_path = write_pdf(pages, tmp_path / "repeat.pdf", COL_WIDTHS, page_height=170.0)
    csv_path = write_csv(rows, tmp_path / "repeat.csv")
    docx_path = write_docx(rows, tmp_path / "repeat.docx")

    assert pypdf.PdfReader(str(pdf_path)).pages

    with csv_path.open(newline="") as fh:
        loaded = list(csv_mod.reader(fh))
    assert loaded.count(rows[0].cells) == 1

    doc = Document(str(docx_path))
    header_rows = [
        [cell.text for cell in row.cells]
        for row in doc.tables[0].rows
        if [cell.text for cell in row.cells] == rows[0].cells
    ]
    assert len(header_rows) == 1
